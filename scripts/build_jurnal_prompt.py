import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_jurnal_prompt_doc():
    doc = docx.Document()

    # A4 Page setup
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

    COLOR_PRIMARY = RGBColor(0x1E, 0x1B, 0x4B)
    COLOR_SECONDARY = RGBColor(0x31, 0x2E, 0x81)
    COLOR_ACCENT = RGBColor(0x43, 0x38, 0xCA)
    COLOR_TEXT = RGBColor(0x1F, 0x29, 0x37)
    COLOR_MUTED = RGBColor(0x4B, 0x55, 0x63)

    def add_p(text, bold_prefix=None, italic=False, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6):
        p = doc.add_paragraph()
        p.alignment = align
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.2
        if bold_prefix:
            r_b = p.add_run(bold_prefix)
            r_b.bold = True
            r_b.font.size = Pt(10.5)
            r_b.font.color.rgb = COLOR_TEXT
        r = p.add_run(text)
        r.italic = italic
        r.font.size = Pt(10.5)
        r.font.color.rgb = COLOR_TEXT
        return p

    def add_heading1(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(8)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(15)
        r.font.color.rgb = COLOR_PRIMARY
        return p

    def add_heading2(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(12)
        r.font.color.rgb = COLOR_SECONDARY
        return p

    def add_prompt_box(prompt_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.27)
        cell = tbl.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:left w:val="single" w:sz="24" w:space="0" w:color="0891B2"/>
                <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC" w:val="clear"/>')
        tcPr.append(shd)
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="140" w:type="dxa"/>
                <w:left w:w="200" w:type="dxa"/>
                <w:bottom w:w="140" w:type="dxa"/>
                <w:right w:w="160" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(prompt_text)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    def add_screenshot_box(instruction_text):
        tbl = doc.add_table(rows=1, cols=1)
        tbl.autofit = False
        tbl.columns[0].width = Inches(6.27)
        cell = tbl.cell(0, 0)
        tcPr = cell._tc.get_or_add_tcPr()
        borders = parse_xml(f'''
            <w:tcBorders {nsdecls("w")}>
                <w:top w:val="dashed" w:sz="8" w:space="0" w:color="9A8AD0"/>
                <w:left w:val="dashed" w:sz="8" w:space="0" w:color="9A8AD0"/>
                <w:bottom w:val="dashed" w:sz="8" w:space="0" w:color="9A8AD0"/>
                <w:right w:val="dashed" w:sz="8" w:space="0" w:color="9A8AD0"/>
            </w:tcBorders>
        ''')
        tcPr.append(borders)
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAF8FF" w:val="clear"/>')
        tcPr.append(shd)
        tcMar = parse_xml(f'''
            <w:tcMar {nsdecls("w")}>
                <w:top w:w="120" w:type="dxa"/>
                <w:left w:w="200" w:type="dxa"/>
                <w:bottom w:w="120" w:type="dxa"/>
                <w:right w:w="200" w:type="dxa"/>
            </w:tcMar>
        ''')
        tcPr.append(tcMar)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(3)
        r1 = p.add_run("[ Tempat screenshot ]")
        r1.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = RGBColor(0x43, 0x38, 0xCA)
        p2 = cell.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_after = Pt(0)
        r2 = p2.add_run(instruction_text)
        r2.italic = True
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    def add_styled_table(headers, rows_data, col_widths=None):
        table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
        table.autofit = False
        for col_idx, h_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.text = h_text
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
        for row_idx, r_data in enumerate(rows_data):
            for col_idx, val in enumerate(r_data):
                cell = table.cell(row_idx + 1, col_idx)
                cell.text = str(val)
                if col_widths and col_idx < len(col_widths):
                    cell.width = col_widths[col_idx]
        for r_idx, row in enumerate(table.rows):
            is_h = (r_idx == 0)
            for c_idx, cell in enumerate(row.cells):
                tcPr = cell._tc.get_or_add_tcPr()
                borders = parse_xml(f'''
                    <w:tcBorders {nsdecls("w")}>
                        <w:top w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:left w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                        <w:right w:val="single" w:sz="4" w:space="0" w:color="CBD5E1"/>
                    </w:tcBorders>
                ''')
                tcPr.append(borders)
                if is_h:
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="EEF2FF" w:val="clear"/>')
                else:
                    bg = "FFFFFF" if r_idx % 2 == 1 else "F8FAFC"
                    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}" w:val="clear"/>')
                tcPr.append(shd)
                tcMar = parse_xml(f'''
                    <w:tcMar {nsdecls("w")}>
                        <w:top w:w="90" w:type="dxa"/>
                        <w:left w:w="130" w:type="dxa"/>
                        <w:bottom w:w="90" w:type="dxa"/>
                        <w:right w:w="130" w:type="dxa"/>
                    </w:tcMar>
                ''')
                tcPr.append(tcMar)
                vAlign = parse_xml(f'<w:vAlign {nsdecls("w")} w:val="center"/>')
                tcPr.append(vAlign)
                p = cell.paragraphs[0]
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for run in p.runs:
                    if is_h:
                        run.bold = True
                        run.font.size = Pt(9.5)
                        run.font.color.rgb = COLOR_PRIMARY
                    else:
                        run.font.size = Pt(9)
                        run.font.color.rgb = COLOR_TEXT
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(6)

    # Header / Title
    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(4)
    r_sub = p_sub.add_run("DOKUMEN JURNAL PROMPT AI AGENT (MAKSIMAL 5 PROMPT)")
    r_sub.bold = True
    r_sub.font.size = Pt(12)
    r_sub.font.color.rgb = COLOR_MUTED

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("SpellCode")
    r_title.bold = True
    r_title.font.size = Pt(32)
    r_title.font.color.rgb = COLOR_PRIMARY

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(20)
    r_m = p_meta.add_run("Peserta: Firman Dwi Nugraha | No. Peserta: 27 | Website: https://www.spellcode.my.id/ | Repo: https://github.com/FirmanDN0/SpellCode")
    r_m.font.size = Pt(9.5)
    r_m.font.color.rgb = COLOR_MUTED

    # Overview
    add_heading1("1. Ketentuan dan Metodologi Jurnal Prompt")
    add_p("Jurnal ini mendokumentasikan 5 prompt utama yang digunakan dalam memandu AI Agent (Antigravity IDE) untuk merancang, mengimplementasikan, menguji, dan mendeploy SpellCode. Setiap prompt disertai dengan teks prompt asli, sasaran teknis, dan verifikasi hasil aktual berbasis automated tests.")

    # Prompt 01
    add_heading1("2. Jurnal 5 Prompt Utama")
    
    add_heading2("Prompt 01: Perancangan Arsitektur Bersih & Mesin Virtual AST Deterministik")
    add_p("Sasaran: Membangun fondasi engine simulasi logika bebas efek samping, Abstract Syntax Tree interpreter, pure state reducer, dan Web Audio API synthesizer.")
    add_prompt_box("""Bertindak sebagai Principal Software Architect. 
Rancang dan bangun fondasi sistem simulasi logika SpellCode:
1. Buat Abstract Syntax Tree (AST) Interpreter murni di lib/engine/interpreter.ts yang mengeksekusi instruksi Move, TurnLeft, TurnRight, Loop, dan Branch secara deterministik.
2. Buat pure state reducer di lib/engine/reducer.ts yang mengelola posisi grid (x, y, arah), inventaris gem, status gembok/gerbang, dan kondisi kemenangan.
3. Buat Web Audio API synthesizer di lib/audio/synth.ts untuk efek audio langkah, petikan gem, mantra, dan victory fanfare tanpa aset file eksternal.
4. Buat unit test lengkap di tests/engine.test.ts untuk memvalidasi seluruh fungsi interpreter. Jalankan npm test dan pastikan semua lulus 100%.""")
    add_p("Hasil Verifikasi: AST Interpreter deterministik selesai, Web Audio synthesizer berfungsi, 32 unit test lulus 100%.")
    add_screenshot_box("Tempel screenshot Antigravity yang memperlihatkan teks prompt asli Tahap 1 dan respons AI Agent.")

    # Prompt 02
    add_heading2("Prompt 02: Pengembangan Antarmuka Meja Mantra & Arena Grid 2D")
    add_p("Sasaran: Membangun Meja Mantra visual (Spell Deck), rak rune (Block Tray), dan visualisator arena grid 2D interaktif dengan tema Friendly Magic Academy.")
    add_prompt_box("""Lanjutkan pengembangan SpellCode berdasarkan arsitektur yang sudah dibangun.
Bangun antarmuka visual gameplay ramah anak di app/adventure/[worldId]/[levelId]/page.tsx:
1. Bangun Meja Mantra (SpellDeck) tempat siswa menyusun balok logika visual dengan kontrol hapus, geser, dan parameter loop/kondisi.
2. Bangun Rak Mantra (BlockTray) yang menampilkan daftar rune yang tersedia sesuai konfigurasi level.
3. Bangun Arena Grid 2D yang merender avatar penyihir, dinding batu, duri rintangan, gem kristal, dan gerbang portal sihir.
4. Terapkan design token bertema 'Friendly Magic Academy' dengan palet warna violet-indigo, font modern, dan responsivitas penuh pada desktop dan tablet.
5. Jalankan verifikasi lint dan typecheck untuk memastikan nol error.""")
    add_p("Hasil Verifikasi: Antarmuka SpellDeck, BlockTray, dan Arena Grid 2D selesai terintegrasi, 0 lint error, 0 type error.")
    add_screenshot_box("Tempel screenshot Antigravity yang memperlihatkan prompt asli Tahap 2, respons AI Agent, atau hasil tampilan antarmuka.")

    # Prompt 03
    add_heading2("Prompt 03: Kurikulum 18 Level & Evaluasi Par Efisiensi Kode")
    add_p("Sasaran: Merancang kurikulum 18 level di 3 dunia petualangan, algoritma Par Calculator 1-3 bintang, serta Step Debugger visual.")
    add_prompt_box("""Bangun kurikulum lengkap dan sistem evaluasi kualitas kode untuk SpellCode:
1. Daftarkan 18 level terstruktur dalam lib/levels/ (World 1 Whispering Woods 6 level sekuensial, World 2 Crystal Caverns 6 level loops, World 3 Cloud Citadel 6 level kondisional).
2. Terapkan sistem Par Calculator di lib/engine/parEvaluator.ts untuk mengevaluasi efisiensi susunan blok kode dan memberikan 1, 2, atau 3 bintang secara objektif.
3. Buat fitur Step Debugger di komponen arena untuk mengeksekusi instruksi langkah demi langkah dengan highlight baris aktif.
4. Buat unit test komprehensif di tests/levels.test.ts dan tests/parEvaluator.test.ts. Jalankan vitest dan verifikasi seluruh kurikulum valid.""")
    add_p("Hasil Verifikasi: 18 level terdaftar lengkap, sistem Par bintang 1-3 aktif, Step Debugger berfungsi presisi.")
    add_screenshot_box("Tempel screenshot Antigravity yang memperlihatkan prompt kurikulum Tahap 3 dan respons hasil pengujian otomatis.")

    # Prompt 04
    add_heading2("Prompt 04: Arsitektur Progres Ganda (Local-First + Supabase Cloud) & Keamanan")
    add_p("Sasaran: Mengimplementasikan sistem penyimpanan ganda offline-first (localStorage) dengan sinkronisasi bi-directional ke Supabase Cloud serta dialog gembok ramah anak.")
    add_prompt_box("""Terapkan arsitektur penyimpanan progres ganda dan keamanan untuk SpellCode:
1. Bangun sistem Local-First di lib/storage/progressStore.ts yang menyimpan progres ke localStorage seketika tanpa jeda server.
2. Buat integrasi Supabase di lib/supabase/client.ts dan lib/supabase/server.ts dengan skema database profiles dan level_progress.
3. Implementasikan fungsi merge bi-directional yang otomatis menyinkronkan progres lokal ke cloud saat pengguna login.
4. Buat Dialog Gembok Ramah Anak di komponen World Map saat pemain mengakses level yang belum terbuka, mencegah akses ilegal tanpa menampilkan pesan error yang menakutkan.
5. Buat tes integrasi dan jalankan npm run build untuk memverifikasi keamanan prerender.""")
    add_p("Hasil Verifikasi: Penyimpanan offline-first lancar, Supabase cloud sync aktif, dialog gembok aman, build Next.js lulus tanpa error prerender.")
    add_screenshot_box("Tempel screenshot Antigravity yang memperlihatkan prompt Tahap 4 dan hasil implementasi sinkronisasi progres.")

    # Prompt 05
    add_heading2("Prompt 05: Audit Kualitas Menyeluruh, Aksesibilitas WCAG, & Deployment Vercel")
    add_p("Sasaran: Audit menyeluruh bug dan race conditions, pengujian aksesibilitas WCAG 2.1 AA, automated E2E testing Playwright, dan deployment live ke Vercel.")
    add_prompt_box("""Lakukan audit kualitas komprehensif, pengujian end-to-end, dan persiapan deployment SpellCode:
1. Lakukan audit kode menyeluruh: perbaiki seluruh type error, lint warning, dan race condition pada eksekusi tombol Run berulang.
2. Pastikan aksesibilitas memenuhi standar WCAG 2.1 AA (kontras rasio warna > 4.5:1, focus-visible ring pada tombol, dan reduced-motion support).
3. Buat dan jalankan automated E2E test menggunakan Playwright yang menguji alur beranda, navigasi peta, arena gameplay, step debugger, dan dialog gembok.
4. Buat panduan deployment docs/DEPLOYMENT_GUIDE.md untuk Vercel dan Supabase.
5. Jalankan verifikasi final: npm run typecheck, npm run lint, npm test, dan npx playwright test. Laporkan seluruh bukti pengujian.""")
    add_p("Hasil Verifikasi: 173 Vitest test passed (100%), 13 Playwright E2E test passed (100%), 0 lint/type error, website live aktif di Vercel.")
    add_screenshot_box("Tempel screenshot Antigravity asli yang memperlihatkan prompt Tahap 5 dan respons hasil audit akhir.")

    # Ringkasan tabel
    add_heading1("3. Ringkasan Eksekusi dan Verifikasi Metrik")
    headers_sum = ["No", "Prompt / Tahap", "Metrik & Hasil Pengujian Aktual", "Status"]
    rows_sum = [
        ["1", "Prompt 01: Arsitektur AST VM", "Deterministic interpreter, Web Audio synth, 32 unit test passed", "Lulus 100%"],
        ["2", "Prompt 02: Antarmuka SpellDeck & Grid", "Meja mantra, rak rune, arena grid 2D, 0 type/lint errors", "Lulus 100%"],
        ["3", "Prompt 03: Kurikulum 18 Level & Par", "18 level terdaftar, Par evaluation 1-3 stars, Step Debugger", "Lulus 100%"],
        ["4", "Prompt 04: Local-First & Supabase", "Bi-directional merge, child-friendly locked dialog, safe prerender", "Lulus 100%"],
        ["5", "Prompt 05: Quality Audit & Deployment", "173 Vitest tests passed, 13 Playwright E2E tests passed, Live Vercel", "Lulus 100%"]
    ]
    add_styled_table(headers_sum, rows_sum, [Inches(0.4), Inches(2.2), Inches(2.8), Inches(0.87)])

    output_path = "SpellCode_Jurnal_Prompt.docx"
    doc.save(output_path)
    print(f"File created: {output_path}")

if __name__ == "__main__":
    create_jurnal_prompt_doc()
