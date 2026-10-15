import os
import re
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

COLOR_PRIMARY = RGBColor(0x1E, 0x1B, 0x4B)   # Deep Indigo #1E1B4B
COLOR_SECONDARY = RGBColor(0x31, 0x2E, 0x81) # Indigo #312E81
COLOR_ACCENT = RGBColor(0x43, 0x38, 0xCA)    # Purple-Indigo #4338CA
COLOR_TEXT = RGBColor(0x1F, 0x29, 0x37)      # Slate Dark #1F2937
COLOR_MUTED = RGBColor(0x4B, 0x55, 0x63)     # Slate Muted #4B5563
COLOR_BORDER = "D1D5DB"

def setup_page(doc):
    for sec in doc.sections:
        sec.page_width = Inches(8.27)
        sec.page_height = Inches(11.69)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(0.9)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)

def add_formatted_runs(p, text, base_size=10.5, base_color=COLOR_TEXT):
    pattern = re.compile(r'(\*\*.*?\*\*|`.*?`|\*.*?\*|\[.*?\]\(.*?\))')
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith('**') and part.endswith('**') and len(part) >= 4:
            r = p.add_run(part[2:-2])
            r.bold = True
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color
        elif part.startswith('`') and part.endswith('`') and len(part) >= 2:
            r = p.add_run(part[1:-1])
            r.font.name = 'Consolas'
            r.font.size = Pt(base_size - 0.5)
            r.font.color.rgb = RGBColor(0x6D, 0x28, 0xD9)
        elif part.startswith('*') and part.endswith('*') and len(part) >= 2:
            r = p.add_run(part[1:-1])
            r.italic = True
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color
        elif part.startswith('[') and '](' in part and part.endswith(')'):
            link_match = re.match(r'\[(.*?)\]\((.*?)\)', part)
            if link_match:
                link_text = link_match.group(1)
                r = p.add_run(link_text)
                r.font.size = Pt(base_size)
                r.font.color.rgb = RGBColor(0x25, 0x63, 0xEB)
                r.underline = True
            else:
                r = p.add_run(part)
                r.font.size = Pt(base_size)
                r.font.color.rgb = base_color
        else:
            r = p.add_run(part)
            r.font.size = Pt(base_size)
            r.font.color.rgb = base_color

def create_styled_table(doc, headers, rows):
    col_count = len(headers)
    if col_count == 0:
        return
    tbl = doc.add_table(rows=len(rows) + 1, cols=col_count)
    tbl.autofit = True
    
    # Header row
    hdr_cells = tbl.rows[0].cells
    for i, h in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        add_formatted_runs(p, h.strip(), base_size=9.5, base_color=RGBColor(0xFF, 0xFF, 0xFF))
        for r in p.runs:
            r.bold = True
        tcPr = hdr_cells[i]._tc.get_or_add_tcPr()
        shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1E1B4B" w:val="clear"/>')
        tcPr.append(shd)

    # Data rows
    for r_idx, row in enumerate(rows):
        row_cells = tbl.rows[r_idx + 1].cells
        bg_fill = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx in range(min(col_count, len(row))):
            row_cells[c_idx].text = ""
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_runs(p, row[c_idx].strip(), base_size=9.0, base_color=COLOR_TEXT)
            tcPr = row_cells[c_idx]._tc.get_or_add_tcPr()
            shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg_fill}" w:val="clear"/>')
            tcPr.append(shd)

    # Borders for all cells
    for row in tbl.rows:
        for cell in row.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            borders = parse_xml(f'''
                <w:tcBorders {nsdecls("w")}>
                    <w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>
                    <w:left w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>
                    <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>
                    <w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_BORDER}"/>
                </w:tcBorders>
            ''')
            tcPr.append(borders)
            tcMar = parse_xml(f'''
                <w:tcMar {nsdecls("w")}>
                    <w:top w:w="80" w:type="dxa"/>
                    <w:left w:w="120" w:type="dxa"/>
                    <w:bottom w:w="80" w:type="dxa"/>
                    <w:right w:w="120" w:type="dxa"/>
                </w:tcMar>
            ''')
            tcPr.append(tcMar)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(6)

def add_callout(doc, quote_lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.27)
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none"/>
            <w:left w:val="single" w:sz="24" w:space="0" w:color="4338CA"/>
            <w:bottom w:val="none"/>
            <w:right w:val="none"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F3FF" w:val="clear"/>')
    tcPr.append(shd)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="100" w:type="dxa"/>
            <w:left w:w="180" w:type="dxa"/>
            <w:bottom w:w="100" w:type="dxa"/>
            <w:right w:w="140" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

    for idx, ql in enumerate(quote_lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(3)
        add_formatted_runs(p, ql, base_size=9.5, base_color=COLOR_TEXT)
        for r in p.runs:
            r.italic = True

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def add_code_block(doc, code_lines):
    tbl = doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.27)
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>
            <w:left w:val="single" w:sz="18" w:space="0" w:color="6D28D9"/>
            <w:bottom w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>
            <w:right w:val="single" w:sz="6" w:space="0" w:color="E2E8F0"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F8FAFC" w:val="clear"/>')
    tcPr.append(shd)
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="80" w:type="dxa"/>
            <w:left w:w="140" w:type="dxa"/>
            <w:bottom w:w="80" w:type="dxa"/>
            <w:right w:w="140" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

    for idx, cl in enumerate(code_lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(cl)
        r.font.name = 'Consolas'
        r.font.size = Pt(8.5)
        r.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    sp = doc.add_paragraph()
    sp.paragraph_format.space_after = Pt(4)

def convert_markdown_to_docx(md_path, docx_path, doc_title=None):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    doc = docx.Document()
    setup_page(doc)

    in_code_block = False
    code_buffer = []
    
    in_table = False
    table_headers = []
    table_rows = []

    in_quote = False
    quote_buffer = []

    idx = 0
    total = len(lines)

    while idx < total:
        raw_line = lines[idx]
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        # Handle Code Block ```
        if stripped.startswith('```'):
            if in_code_block:
                add_code_block(doc, code_buffer)
                code_buffer = []
                in_code_block = False
            else:
                # Flush quote/table if any
                if in_quote:
                    add_callout(doc, quote_buffer)
                    quote_buffer = []
                    in_quote = False
                if in_table:
                    create_styled_table(doc, table_headers, table_rows)
                    table_headers, table_rows = [], []
                    in_table = False
                in_code_block = True
            idx += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            idx += 1
            continue

        # Handle Blockquotes >
        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            quote_buffer.append(quote_text)
            in_quote = True
            idx += 1
            continue
        elif in_quote:
            add_callout(doc, quote_buffer)
            quote_buffer = []
            in_quote = False

        # Handle Tables | ... |
        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            # Check if separator row like |---|---|
            is_sep = all(re.match(r'^:?-+:?$', c) for c in cells) if cells else False
            if not in_table:
                table_headers = cells
                table_rows = []
                in_table = True
            elif is_sep:
                pass  # skip separator
            else:
                table_rows.append(cells)
            idx += 1
            continue
        elif in_table:
            create_styled_table(doc, table_headers, table_rows)
            table_headers, table_rows = [], []
            in_table = False

        # Empty lines
        if not stripped:
            idx += 1
            continue

        # Horizontal Rule ---
        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            # Subtle divider
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("_________________________________________________________________________________")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            idx += 1
            continue

        # Headings
        if stripped.startswith('# '):
            # Document Main Title
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title_text)
            r.bold = True
            r.font.size = Pt(20)
            r.font.color.rgb = COLOR_PRIMARY

            # Add Metadata Subtitle
            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(16)
            r_sub = p_sub.add_run("SpellCode Platform | Peserta: Firman Dwi Nugraha (No. 27)\nWebsite: https://www.spellcode.my.id/ | GitHub: https://github.com/FirmanDN0/SpellCode")
            r_sub.font.size = Pt(9.5)
            r_sub.font.color.rgb = COLOR_MUTED

            idx += 1
            continue

        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = COLOR_PRIMARY
            idx += 1
            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = COLOR_SECONDARY
            idx += 1
            continue

        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_ACCENT
            idx += 1
            continue

        # Unordered Bullet List (* or -)
        if re.match(r'^[\*\-]\s+', stripped):
            item_text = re.sub(r'^[\*\-]\s+', '', stripped)
            indent_level = 0
            # Check indentation in raw_line
            leading_spaces = len(raw_line) - len(raw_line.lstrip(' '))
            if leading_spaces >= 4:
                indent_level = 1
            elif leading_spaces >= 2:
                indent_level = 0.5

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25 + indent_level * 0.25)
            
            bullet_sym = "◦ " if indent_level > 0 else "• "
            r_b = p.add_run(bullet_sym)
            r_b.bold = True
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = COLOR_ACCENT

            add_formatted_runs(p, item_text, base_size=10.0, base_color=COLOR_TEXT)
            idx += 1
            continue

        # Ordered Numbered List (1. ...)
        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            num_str = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25)
            r_n = p.add_run(f"{num_str}. ")
            r_n.bold = True
            r_n.font.size = Pt(10)
            r_n.font.color.rgb = COLOR_PRIMARY

            add_formatted_runs(p, item_text, base_size=10.0, base_color=COLOR_TEXT)
            idx += 1
            continue

        # Regular Body Paragraph
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_runs(p, stripped, base_size=10.5, base_color=COLOR_TEXT)
        idx += 1

    # Final flushes
    if in_code_block and code_buffer:
        add_code_block(doc, code_buffer)
    if in_quote and quote_buffer:
        add_callout(doc, quote_buffer)
    if in_table and table_headers:
        create_styled_table(doc, table_headers, table_rows)

    doc.save(docx_path)
    print(f"Berhasil membuat: {docx_path}")

def parse_markdown_into_doc(doc, md_path):
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    code_buffer = []
    
    in_table = False
    table_headers = []
    table_rows = []

    in_quote = False
    quote_buffer = []

    idx = 0
    total = len(lines)

    while idx < total:
        raw_line = lines[idx]
        line = raw_line.rstrip('\r\n')
        stripped = line.strip()

        if stripped.startswith('```'):
            if in_code_block:
                add_code_block(doc, code_buffer)
                code_buffer = []
                in_code_block = False
            else:
                if in_quote:
                    add_callout(doc, quote_buffer)
                    quote_buffer = []
                    in_quote = False
                if in_table:
                    create_styled_table(doc, table_headers, table_rows)
                    table_headers, table_rows = [], []
                    in_table = False
                in_code_block = True
            idx += 1
            continue

        if in_code_block:
            code_buffer.append(line)
            idx += 1
            continue

        if stripped.startswith('>'):
            quote_text = stripped.lstrip('>').strip()
            quote_buffer.append(quote_text)
            in_quote = True
            idx += 1
            continue
        elif in_quote:
            add_callout(doc, quote_buffer)
            quote_buffer = []
            in_quote = False

        if stripped.startswith('|') and stripped.endswith('|'):
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            is_sep = all(re.match(r'^:?-+:?$', c) for c in cells) if cells else False
            if not in_table:
                table_headers = cells
                table_rows = []
                in_table = True
            elif is_sep:
                pass
            else:
                table_rows.append(cells)
            idx += 1
            continue
        elif in_table:
            create_styled_table(doc, table_headers, table_rows)
            table_headers, table_rows = [], []
            in_table = False

        if not stripped:
            idx += 1
            continue

        if re.match(r'^-{3,}$', stripped) or re.match(r'^\*{3,}$', stripped):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run("_________________________________________________________________________________")
            r.font.size = Pt(8)
            r.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)
            idx += 1
            continue

        if stripped.startswith('# '):
            title_text = stripped[2:].strip()
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(title_text)
            r.bold = True
            r.font.size = Pt(18)
            r.font.color.rgb = COLOR_PRIMARY

            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(14)
            r_sub = p_sub.add_run("SpellCode Platform | Peserta: Firman Dwi Nugraha (No. 27)\nWebsite: https://www.spellcode.my.id/ | GitHub: https://github.com/FirmanDN0/SpellCode")
            r_sub.font.size = Pt(9.5)
            r_sub.font.color.rgb = COLOR_MUTED

            idx += 1
            continue

        if stripped.startswith('## '):
            h_text = stripped[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(14)
            r.font.color.rgb = COLOR_PRIMARY
            idx += 1
            continue

        if stripped.startswith('### '):
            h_text = stripped[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(12)
            r.font.color.rgb = COLOR_SECONDARY
            idx += 1
            continue

        if stripped.startswith('#### '):
            h_text = stripped[5:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            r = p.add_run(h_text)
            r.bold = True
            r.font.size = Pt(10.5)
            r.font.color.rgb = COLOR_ACCENT
            idx += 1
            continue

        if re.match(r'^[\*\-]\s+', stripped):
            item_text = re.sub(r'^[\*\-]\s+', '', stripped)
            indent_level = 0
            leading_spaces = len(raw_line) - len(raw_line.lstrip(' '))
            if leading_spaces >= 4:
                indent_level = 1
            elif leading_spaces >= 2:
                indent_level = 0.5

            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25 + indent_level * 0.25)
            
            bullet_sym = "◦ " if indent_level > 0 else "• "
            r_b = p.add_run(bullet_sym)
            r_b.bold = True
            r_b.font.size = Pt(10)
            r_b.font.color.rgb = COLOR_ACCENT

            add_formatted_runs(p, item_text, base_size=10.0, base_color=COLOR_TEXT)
            idx += 1
            continue

        num_match = re.match(r'^(\d+)\.\s+(.*)', stripped)
        if num_match:
            num_str = num_match.group(1)
            item_text = num_match.group(2)
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.left_indent = Inches(0.25)
            r_n = p.add_run(f"{num_str}. ")
            r_n.bold = True
            r_n.font.size = Pt(10)
            r_n.font.color.rgb = COLOR_PRIMARY

            add_formatted_runs(p, item_text, base_size=10.0, base_color=COLOR_TEXT)
            idx += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        add_formatted_runs(p, stripped, base_size=10.5, base_color=COLOR_TEXT)
        idx += 1

    if in_code_block and code_buffer:
        add_code_block(doc, code_buffer)
    if in_quote and quote_buffer:
        add_callout(doc, quote_buffer)
    if in_table and table_headers:
        create_styled_table(doc, table_headers, table_rows)

def build_all_docs():
    base_dir = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
    docs_dir = os.path.join(base_dir, "docs")

    doc_files = [
        ("PRD.md", "SpellCode_1_PRD.docx"),
        ("PRODUCT_SPEC.md", "SpellCode_2_Product_Spec.docx"),
        ("TECHNICAL_ARCHITECTURE.md", "SpellCode_3_Technical_Architecture.docx"),
        ("DESIGN_SYSTEM.md", "SpellCode_4_Design_System.docx"),
        ("DEPLOYMENT_GUIDE.md", "SpellCode_5_Deployment_Guide.docx")
    ]

    # 1. Build 5 Individual DOCX
    for md_file, docx_file in doc_files:
        md_path = os.path.join(docs_dir, md_file)
        docx_path = os.path.join(docs_dir, docx_file)
        if os.path.exists(md_path):
            doc = docx.Document()
            setup_page(doc)
            parse_markdown_into_doc(doc, md_path)
            doc.save(docx_path)
            print(f"Berhasil membuat: {docx_path}")
        else:
            print(f"Warning: File {md_path} tidak ditemukan!")

    # 2. Build 1 Merged Master DOCX (Bundled)
    master_doc = docx.Document()
    setup_page(master_doc)

    # Master Cover
    p_title = master_doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(40)
    p_title.paragraph_format.space_after = Pt(8)
    r_t = p_title.add_run("SPELLCODE: BUNDEL DOKUMEN ACUAN AI")
    r_t.bold = True
    r_t.font.size = Pt(22)
    r_t.font.color.rgb = COLOR_PRIMARY

    p_sub = master_doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(20)
    r_s = p_sub.add_run("Kompilasi 5 Dokumen Spesifikasi, Arsitektur, Desain, & Deployment Resmi\nPlatform Edukasi Pemrograman Fantasi Anak")
    r_s.italic = True
    r_s.font.size = Pt(11)
    r_s.font.color.rgb = COLOR_SECONDARY

    # Metadata Box
    tbl = master_doc.add_table(rows=1, cols=1)
    tbl.autofit = False
    tbl.columns[0].width = Inches(6.27)
    cell = tbl.cell(0, 0)
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="FAF8FF" w:val="clear"/>')
    tcPr.append(shd)
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="12" w:space="0" w:color="4338CA"/>
            <w:left w:val="single" w:sz="12" w:space="0" w:color="4338CA"/>
            <w:bottom w:val="single" w:sz="12" w:space="0" w:color="4338CA"/>
            <w:right w:val="single" w:sz="12" w:space="0" w:color="4338CA"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)
    p_box = cell.paragraphs[0]
    p_box.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_box.paragraph_format.space_before = Pt(10)
    p_box.paragraph_format.space_after = Pt(10)
    r_box = p_box.add_run("Peserta: Firman Dwi Nugraha | No. Peserta: 27\nDomain Resmi: https://www.spellcode.my.id/\nMirror Vercel: https://spell-code-nine.vercel.app/\nGitHub: https://github.com/FirmanDN0/SpellCode\nTanggal Rilis: Oktober 2026")
    r_box.font.size = Pt(10)
    r_box.font.color.rgb = COLOR_TEXT

    # Daftar Isi Singkat
    p_toc_title = master_doc.add_paragraph()
    p_toc_title.paragraph_format.space_before = Pt(28)
    p_toc_title.paragraph_format.space_after = Pt(8)
    r_tt = p_toc_title.add_run("Daftar Dokumen yang Terkandung:")
    r_tt.bold = True
    r_tt.font.size = Pt(13)
    r_tt.font.color.rgb = COLOR_PRIMARY

    toc_items = [
        ("Dokumen 1: Product Requirements Document (PRD)", "Visi, misi, persona pengguna, kurikulum 3 dunia (18 level), mekanik par."),
        ("Dokumen 2: Product Specification & Game Design", "Detail teknis gameplay loop, 8 instruksi rune sihir, evaluasi bintang, audio-visual."),
        ("Dokumen 3: Technical Architecture Specification", "Arsitektur client VM, AST compiler, deterministic engine, PostgreSQL RLS."),
        ("Dokumen 4: Design System Specification", "Palet warna Whimsical Fantasy, kontras WCAG AA, tipografi ramah anak, ukuran komponen."),
        ("Dokumen 5: Deployment & Operations Guide", "Setup Supabase, environment variables, verifikasi Vercel, dan security checklist.")
    ]
    for doc_name, doc_desc in toc_items:
        p_item = master_doc.add_paragraph()
        p_item.paragraph_format.space_before = Pt(2)
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.left_indent = Inches(0.25)
        r_num = p_item.add_run("• ")
        r_num.bold = True
        r_num.font.color.rgb = COLOR_ACCENT
        r_dn = p_item.add_run(f"{doc_name}\n  ")
        r_dn.bold = True
        r_dn.font.size = Pt(10)
        r_dn.font.color.rgb = COLOR_PRIMARY
        r_dd = p_item.add_run(doc_desc)
        r_dd.font.size = Pt(9.5)
        r_dd.font.color.rgb = COLOR_MUTED

    # Append each of the 5 files
    for idx, (md_file, _) in enumerate(doc_files):
        master_doc.add_page_break()
        md_path = os.path.join(docs_dir, md_file)
        if os.path.exists(md_path):
            parse_markdown_into_doc(master_doc, md_path)

    master_docx_path = os.path.join(docs_dir, "SpellCode_Dokumen_Acuan_AI_Lengkap.docx")
    master_doc.save(master_docx_path)
    print(f"Berhasil membuat bundel lengkap: {master_docx_path}")

if __name__ == "__main__":
    build_all_docs()
